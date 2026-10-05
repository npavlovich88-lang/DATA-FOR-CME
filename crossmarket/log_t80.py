"""Append T80 to out/test_log.xlsx with its verdict block. The log continues the numbering of
the 79-test log in the quant project; that file was not available, so this one starts at T80."""

from __future__ import annotations

import json
import pathlib

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from quant.verdict import Verdict

OUT = pathlib.Path(__file__).parent / "out"
LOG = OUT / "test_log.xlsx"
COLS = ["test", "date", "title", "data", "window", "configs", "headline", "checks", "verdict",
        "binding reason", "next"]


def verdict_t80(r: dict) -> Verdict:
    days = next(iter(r.values()))["base"]["days"]
    fam = [v["perm_p_family"] for v in r.values()]
    pf = f"{min(fam):.2f}-{max(fam):.2f}"
    rows = []
    for k, v in r.items():
        b = v["base"]
        lo, hi = b["sharpe_ci"]
        rows.append(
            f"{k}: Sharpe {b['sharpe']:+.2f} [95% {lo:+.1f},{hi:+.1f}], 3x cost "
            f"{v['cost3x']['sharpe']:+.2f}, ex-top-3-days {b['sharpe_ex_top3']:+.2f}, halves "
            f"{v['halves'][0]:+.2f}/{v['halves'][1]:+.2f}, perm p {v['perm_p']:.2f}, "
            f"family p {v['perm_p_family']:.2f}"
        )
    return Verdict.inconclusive(
        "T80 4h trend following (breakout + EMA), 12 markets, train",
        reason=f"{days} scored days cannot resolve a trend Sharpe: every 95% interval spans "
        "roughly -2 to +2.4, so the published 0.3-0.8 and zero are equally consistent with the "
        f"data. No config survives the best-of-4 permutation null (family p {pf}).",
        evidence=rows
        + [
            "Positive configs are carried by ES/NQ/RTY/YM, one correlated long-equity bet over "
            "2025 (standing check 21); rates, HG and 6J lose under all four configs.",
            "Every config is negative once its best three days are dropped (check 17), and three "
            "of four flip sign between halves (check 1).",
        ],
        caveat="Yahoo 1h mixes contracts around rolls; cleaned by calendar-window exclusion. The "
        "strict cleaning variant moves Sharpe by <0.35 and changes no conclusion. Realised vol "
        "~16% against a 10% target: the four equity indices are one risk, not four.",
        next_step="Run the same four rules on the 26-year daily set (check 22: largest sample "
        "first) before spending anything on the 2-year intraday set.",
    )


def append(test: str, row: list, block: str) -> None:
    """Add one row to the log and the verdict block on its own sheet. Refuses duplicates."""
    if LOG.exists():
        wb = load_workbook(LOG)
        ws = wb["log"]
        if any(c.value == test for c in ws["A"]):
            raise SystemExit(f"{test} already logged")
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "log"
        ws.append(COLS)
        for c in ws[1]:
            c.font = Font(bold=True)
    ws.append(row)
    for c in ws[ws.max_row]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    vs = wb.create_sheet(f"{test} verdict")
    for line in block.splitlines():
        vs.append([line])
    vs.column_dimensions["A"].width = 100
    for c in vs["A"]:
        c.font = Font(name="Consolas")
    wb.save(LOG)


def main() -> None:
    r = json.loads((OUT / "t80_results.json").read_text())
    v = verdict_t80(r)
    block = v.render()
    best = max(r, key=lambda k: r[k]["base"]["sharpe"])
    row = [
        "T80",
        "2026-10-05",
        "4h trend following: Donchian breakout + EMA crossover, inverse-vol, equal risk",
        "Yahoo =F 1h -> session 4h, 12 markets (SB excluded: contract-mixed), roll-cleaned",
        f"train only, scored {r[best]['window'][0]}..{r[best]['window'][1]} "
        f"({r[best]['base']['days']} days)",
        ", ".join(r),
        "; ".join(f"{k} SR {r[k]['base']['sharpe']:+.2f}" for k in r),
        "1,3,7,9,12,16,17,20,21,22 applied; costs 1x/3x; lag 2; strict cleaning",
        v.state,
        v.reason,
        v.next_step,
    ]
    append("T80", row, block)
    print(block)


if __name__ == "__main__":
    main()
